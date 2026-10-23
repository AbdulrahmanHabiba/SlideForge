import io
import json
import zipfile
import shutil
from unittest.mock import AsyncMock

import fitz
import pytest
from PIL import Image
from fastapi.testclient import TestClient
from pptx import Presentation

from app.core import (ConversionError, explicitly_downloadable, slide_image_urls,
                      validate_url, images_to_files, convert_source, Source)
from app.main import app

client = TestClient(app)


def image_data(width=960, height=540, color='#65cdbb'):
    img = Image.new('RGB', (width, height), color)
    dest = io.BytesIO()
    img.save(dest, format='PNG')
    return dest.getvalue()


def pdf_data():
    doc = fitz.open()
    page = doc.new_page(width=960, height=540)
    page.insert_text((60, 70), 'Safety presentation sample')
    raw = doc.tobytes()
    doc.close()
    return raw


def test_health_and_frontend():
    assert client.get('/api/health').json()['max_links'] == 100
    assert client.get('/').status_code == 200
    assert 'SlideForge' in client.get('/').text
    assert client.get('/static/app.js').status_code == 200


def test_validation():
    for url in ['http://example.com/x.pdf', 'https://localhost/a', 'https://127.0.0.1/a',
                'https://192.168.1.10/a', 'https://user:pass@example.com/a',
                'https://example.com:333/a']:
        with pytest.raises(ConversionError): validate_url(url)
    assert validate_url('https://files.example.com/a.pdf') == 'https://files.example.com/a.pdf'


def test_slideshare_permission_fails_closed():
    assert not explicitly_downloadable('{"downloadable":false}')
    assert not explicitly_downloadable('<button>Download</button>')
    assert explicitly_downloadable('{"downloadable":true}')
    assert not explicitly_downloadable('{"downloadable":true,"allowDownload":false}')


def test_slide_image_extraction_quality_and_sequence():
    raw = '{"downloadable":true}' + ''.join(
        f'"https://image.slidesharecdn.com/demo/75/talk-{n}-1024.jpg"'
        f'"https://image.slidesharecdn.com/demo/75/talk-{n}-2048.jpg"' for n in range(1, 4)
    )
    images = slide_image_urls(raw)
    assert len(images) == 3
    assert all('2048' in x for x in images)
    incomplete = raw.replace('talk-2-', 'talk-5-')
    assert not slide_image_urls(incomplete)


def test_images_to_both_formats_and_sd_size():
    images = [image_data(width=2000, height=1200, color='#3456ff'), image_data(width=2000, height=1200, color='#33dd77')]
    result = images_to_files(images, 'hd', 'both')
    assert result['pdf'].startswith(b'%PDF-')
    with fitz.open(stream=result['pdf'], filetype='pdf') as document:
        assert len(document) == 2
    presentation = Presentation(io.BytesIO(result['pptx']))
    assert len(presentation.slides) == 2
    smaller = images_to_files(images, 'sd', 'both')
    assert len(smaller['pptx']) < len(result['pptx'])


def test_pdf_preserves_original_and_converts_to_powerpoint():
    original = pdf_data()
    src = Source('testing', original, '.pdf', 'upload')
    result = convert_source(src, 'hd', 'both')
    assert result['pdf'] == original
    assert len(Presentation(io.BytesIO(result['pptx'])).slides) == 1


@pytest.mark.skipif(shutil.which('libreoffice') is None, reason='LibreOffice is not installed')
def test_pptx_round_trip_to_pdf():
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    dest = io.BytesIO()
    presentation.save(dest)
    result = convert_source(Source('demo', dest.getvalue(), '.pptx', 'upload'), 'sd', 'both')
    assert result['pptx'] == dest.getvalue()
    assert result['pdf'].startswith(b'%PDF-')


def test_batch_zip_end_to_end():
    r = client.post('/api/convert', data={'urls': '[]', 'fmt': 'both', 'quality': 'hd'},
                    files=[('files', ('deck.png', image_data(), 'image/png'))])
    assert r.status_code == 200
    assert 'application/zip' in r.headers['content-type']
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    assert {'deck.pdf', 'deck.pptx', 'manifest.json'} <= set(names)
    assert json.loads(z.read('manifest.json'))['items'][0]['files'][0]['name'] == 'deck.pdf'
    assert len(Presentation(io.BytesIO(z.read('deck.pptx'))).slides) == 1


def test_download_endpoint_returns_public_slide_note(monkeypatch):
    monkeypatch.setattr('app.main.get_source_from_url', AsyncMock(return_value=('Public deck', [image_data()])))
    r = client.post('/api/download', data={'urls': '["https://www.slideshare.net/slideshow/demo/123"]', 'fmt': 'both', 'quality': 'sd'})
    assert r.status_code == 200
    z = zipfile.ZipFile(io.BytesIO(r.content))
    item = json.loads(z.read('manifest.json'))['items'][0]
    assert len(item['files']) == 2
    assert 'not the original' in item['note']


def test_mixed_batch_partial_failure():
    r = client.post('/api/convert', data={'urls': '["not-a-link"]', 'fmt': 'pdf', 'quality': 'sd'},
                    files=[('files', ('good.png', image_data(), 'image/png'))])
    assert r.status_code == 200
    z = zipfile.ZipFile(io.BytesIO(r.content))
    items = json.loads(z.read('manifest.json'))['items']
    assert items[0]['error'] is not None
    assert items[1]['files']


def test_batch_rejects_more_than_100():
    r = client.post('/api/convert', data={'urls': json.dumps(['https://a.example.com/test.pdf'] * 101), 'fmt': 'both', 'quality': 'hd'})
    assert r.status_code == 422


def test_scribd_link_reports_real_limitation():
    r = client.post('/api/convert', data={'urls': '["https://www.scribd.com/document/123/test"]', 'fmt': 'both', 'quality': 'hd'})
    assert r.status_code == 422
    assert 'official' in r.json()['items'][0]['error']


def test_malformed_pdf_does_not_succeed():
    r = client.post('/api/convert', data={'urls': '[]', 'fmt': 'both', 'quality': 'hd'},
                    files=[('files', ('fake.pdf', b'%PDF-broken', 'application/pdf'))])
    assert r.status_code == 422


def test_challenge_uses_browser_and_then_downloads_images(monkeypatch):
    from pathlib import Path
    fixture = (Path(__file__).parent / 'fixtures' / 'hospital-public-viewer.html').read_text()
    async def fetch(url, max_bytes=0):
        if 'www.slideshare.net' in url:
            return b'<title>Client Challenge</title>', 'text/html', url
        return image_data(), 'image/png', url
    render = AsyncMock(return_value=fixture)
    monkeypatch.setattr('app.core.fetch_public', fetch)
    monkeypatch.setattr('app.browser.render_public_slideshare', render)
    r = client.post('/api/download', data={'urls': '["https://www.slideshare.net/slideshow/hospital-fire-safety-250965771/250965771"]', 'fmt': 'pdf', 'quality': 'sd'})
    assert r.status_code == 200
    render.assert_awaited_once()
    with zipfile.ZipFile(io.BytesIO(r.content)) as archive:
        item = json.loads(archive.read('manifest.json'))['items'][0]
        with fitz.open(stream=archive.read(item['files'][0]['name']), filetype='pdf') as pdf:
            assert len(pdf) == 49


def test_private_viewer_is_not_retried_in_browser(monkeypatch):
    fixture = '<script id="__NEXT_DATA__">' + json.dumps({'props': {'pageProps': {'slideshow': {'isPrivate': True}}}}) + '</script>'
    monkeypatch.setattr('app.core.fetch_public', AsyncMock(return_value=(fixture.encode(), 'text/html', 'https://www.slideshare.net/slideshow/private/123')))
    render = AsyncMock()
    monkeypatch.setattr('app.browser.render_public_slideshare', render)
    r = client.post('/api/download', data={'urls': '["https://www.slideshare.net/slideshow/private/123"]', 'fmt': 'pdf', 'quality': 'sd'})
    assert r.status_code == 422
    render.assert_not_awaited()
